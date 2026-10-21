# --Genera data/ejemplo/casos_ficticios.xlsx: un archivo de seguimiento con casos inventados.
# La semilla es fija: correrlo dos veces produce el mismo archivo.

import random
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import config  

SEMILLA = 20261020
SALIDA = config.DATA_DIR / "ejemplo" / "casos_ficticios.xlsx"


ENCABEZADOS = [
    "Tipo", "Usuario que reporta", "Fecha de recepcion", "Sprint", "Descripción incidente",
    "id flujo", "nombre_flujo", "Fabrica", "Masificado", "#HA", "DEV", "Fecha asignado",
    "Fecha compromiso DEV", "Fecha fin DEV", "Documentación Tecnica", "Categoria Causa",
    "Sub_Categoría Causa", "Art. Conoc", "Etapa", "Estado", "#HA", "QA",
    "Fecha compromiso QA", "Fecha Fin QA", "Fecha Solicitud VoBo Cliente",
    "Fecha VoBo Cliente", "Fecha de cierre", "Fecha Salida PRD", "Días para cierre",
    "Observaciones", "Causa/Solucion",
]

FLUJOS = {
    "Solicitud de crédito": "FL-001",
    "Apertura de cuenta": "FL-002",
    "Reclamos": "FL-003",
    "Gestión de garantías": "FL-004",
}

ESTADOS_ABIERTOS = [
    ("1. Recepción", "1.1. Nuevo"),
    ("1. Recepción", "1.3. Estimado"),
    ("2. Desarrollo", "2.1. Asignado"),
    ("2. Desarrollo", "2.2. En Ejecución"),
    ("2. Desarrollo", "2.4. Impedimento DEV"),
    ("2. Desarrollo", "2.5. Desplegado a QA"),
    ("3. Pruebas", "3.1. En QA"),
    ("3. Pruebas", "3.3. Bug QA"),
    ("3. Pruebas", "3.6. Pendiente VoBo Usuario"),
    ("4. Paso a Producción y Entrega", "4.1. Pendiente PDN"),
    ("4. Paso a Producción y Entrega", "4.2. En Producción pte validación usuario"),
]
CERRADO = ("4. Paso a Producción y Entrega", "4.3. Cerrado")
CERRADO_SIN_PRD = ("4. Paso a Producción y Entrega", "4.4. Cerrado Sin Paso a Producción")


ESCENARIOS = [
    {
        "flujo": "Solicitud de crédito",
        "descripcion": "Al radicar la solicitud el caso queda sin asignar y no aparece en ninguna bandeja de analistas.",
        "causa": "La regla de asignación apuntaba a un grupo renombrado en la migración y no encontraba responsables.",
        "solucion": "Se actualizó la regla de asignación con el nombre vigente del grupo y se reasignaron los casos pendientes.",
        "categoria": "Parametrización_de_Accesos_Roles_y_Asignación",
        "subcategoria": "Asignación y bandejas",
        "despliegue": False,
    },
    {
        "flujo": "Apertura de cuenta",
        "descripcion": "Usuarios nuevos del área de operaciones no pueden ingresar al flujo; el sistema muestra 'acceso denegado'.",
        "causa": "Los usuarios no fueron agregados al grupo de seguridad del flujo al crear sus cuentas.",
        "solucion": "Se agregaron los usuarios al grupo correspondiente y se incluyó el paso en el procedimiento de alta.",
        "categoria": "Parametrización_de_Accesos_Roles_y_Asignación",
        "subcategoria": "Accesos y usuarios",
        "despliegue": False,
    },
    {
        "flujo": "Gestión de garantías",
        "descripcion": "El módulo administrativo no permite editar el catálogo de tipos de garantía.",
        "causa": "El grupo de administradores funcionales no tenía permiso de edición sobre el catálogo.",
        "solucion": "Se otorgó el permiso de edición al grupo de administradores funcionales.",
        "categoria": "Parametrización_de_Accesos_Roles_y_Asignación",
        "subcategoria": "Grupos, catálogo y módulo administrativo",
        "despliegue": False,
    },
    {
        "flujo": "Reclamos",
        "descripcion": "El campo 'Número de producto' acepta letras y se están radicando reclamos con datos inválidos.",
        "causa": "La validación del campo se perdió al migrar el formulario y quedó como texto libre.",
        "solucion": "Se agregó la validación numérica y de longitud al campo en el formulario.",
        "categoria": "Parametrización_Campos_Reglas_e_Interfaces",
        "subcategoria": "Campos, listas y validaciones",
        "despliegue": True,
    },
    {
        "flujo": "Reclamos",
        "descripcion": "Los casos se marcan como vencidos un día antes de lo que indica el ANS.",
        "causa": "El cálculo del ANS contaba días calendario en lugar de días hábiles.",
        "solucion": "Se ajustó la regla del ANS para usar el calendario de días hábiles.",
        "categoria": "Parametrización_Campos_Reglas_e_Interfaces",
        "subcategoria": "ANS y tiempos",
        "despliegue": True,
    },
    {
        "flujo": "Solicitud de crédito",
        "descripcion": "La ayuda del campo 'Ingresos mensuales' dice que es opcional, pero el sistema no deja avanzar sin él.",
        "causa": "El texto de ayuda no se actualizó cuando el campo pasó a ser obligatorio.",
        "solucion": "Se corrigió el texto de ayuda del campo.",
        "categoria": "Parametrización_Campos_Reglas_e_Interfaces",
        "subcategoria": "Textos, etiquetas y mensajes de ayuda",
        "despliegue": False,
    },
    {
        "flujo": "Gestión de garantías",
        "descripcion": "El informe detallado no muestra la columna de fecha de aprobación.",
        "causa": "La consulta del informe no incluía la fecha de aprobación después del cambio en el modelo de datos.",
        "solucion": "Se agregó el campo a la consulta y a la plantilla del informe.",
        "categoria": "Configuración_Informes",
        "subcategoria": "Contenido del informe detallado",
        "despliegue": True,
    },
    {
        "flujo": "Reclamos",
        "descripcion": "Los supervisores no pueden descargar el informe en Excel; el botón no aparece.",
        "causa": "El permiso de descarga estaba asignado sólo al rol de administrador.",
        "solucion": "Se otorgó el permiso de descarga al rol de supervisor.",
        "categoria": "Configuración_Informes",
        "subcategoria": "Acceso y descarga de informes",
        "despliegue": False,
    },
    {
        "flujo": "Apertura de cuenta",
        "descripcion": "El historial del caso no muestra quién aprobó la solicitud.",
        "causa": "La actividad de aprobación no registraba el usuario en la tabla de trazabilidad.",
        "solucion": "Se modificó la actividad para guardar el usuario y la fecha de aprobación en el historial.",
        "categoria": "Configuración_Informes",
        "subcategoria": "Seguimiento, historial y trazabilidad",
        "despliegue": True,
    },
    {
        "flujo": "Solicitud de crédito",
        "descripcion": "Al radicar desde el canal web aparece de forma intermitente el error 'No se pudo crear el caso'.",
        "causa": "El servicio de numeración devolvía consecutivos duplicados cuando llegaban solicitudes simultáneas.",
        "solucion": "Se cambió la generación del consecutivo a una secuencia de base de datos.",
        "categoria": "Desarrollo",
        "subcategoria": "Radicación y creación de casos",
        "despliegue": True,
    },
    {
        "flujo": "Gestión de garantías",
        "descripcion": "El caso queda detenido en la actividad 'Revisión jurídica' y no avanza aunque se complete el formulario.",
        "causa": "La compuerta de salida evaluaba una variable que nunca se asignaba en ese camino del flujo.",
        "solucion": "Se asignó la variable antes de la compuerta y se reprocesaron los casos detenidos.",
        "categoria": "Desarrollo",
        "subcategoria": "Ciclo del caso y actividades",
        "despliegue": True,
    },
    {
        "flujo": "Reclamos",
        "descripcion": "Los datos de la pestaña 'Detalle' se pierden al guardar el borrador y volver a abrir el caso.",
        "causa": "La acción 'Guardar borrador' no persistía la pestaña de detalle en el registro del caso.",
        "solucion": "Se agregó el guardado de la pestaña de detalle en la acción de borrador.",
        "categoria": "Desarrollo",
        "subcategoria": "Visualización y persistencia de datos",
        "despliegue": True,
    },
    {
        "flujo": "Apertura de cuenta",
        "descripcion": "Alerta de monitoreo: el servicio de validación de identidad responde con error 500 en producción.",
        "causa": "El certificado del servicio externo venció y la integración rechazaba la conexión.",
        "solucion": "Se renovó el certificado y se programó una alerta antes del próximo vencimiento.",
        "categoria": "Desarrollo",
        "subcategoria": "Errores de servicios en producción (monitoreo)",
        "despliegue": False,
    },
    {
        "flujo": "Solicitud de crédito",
        "descripcion": "El correo de aprobación le llega al asesor anterior y no al asesor actual del cliente.",
        "causa": "La notificación tomaba como destinatario al primer asesor asignado y no al vigente.",
        "solucion": "Se cambió la regla para tomar el asesor vigente al momento del envío.",
        "categoria": "Configuración_Notificaciones",
        "subcategoria": "Destinatarios",
        "despliegue": True,
    },
    {
        "flujo": "Reclamos",
        "descripcion": "Los clientes no reciben el correo de confirmación de radicación.",
        "causa": "La plantilla del correo quedó inactiva después de la migración.",
        "solucion": "Se activó la plantilla y se reenviaron las confirmaciones pendientes.",
        "categoria": "Configuración_Notificaciones",
        "subcategoria": "Notificaciones que no llegan",
        "despliegue": False,
    },
    {
        "flujo": "Gestión de garantías",
        "descripcion": "El correo de tarea asignada no muestra el número del caso ni el nombre del cliente.",
        "causa": "La plantilla del correo usaba variables con nombres anteriores a la migración.",
        "solucion": "Se actualizaron las variables de la plantilla con los nombres vigentes.",
        "categoria": "Configuración_Notificaciones",
        "subcategoria": "Contenido del correo",
        "despliegue": False,
    },
    {
        "flujo": "Gestión de garantías",
        "descripcion": "Los responsables reciben dos veces la misma notificación de tarea asignada.",
        "causa": "La notificación estaba configurada tanto en la actividad como en el proceso padre.",
        "solucion": "Se eliminó la notificación duplicada del proceso padre.",
        "categoria": "Configuración_Notificaciones",
        "subcategoria": "Duplicadas o con retraso",
        "despliegue": True,
    },
    {
        "flujo": "Apertura de cuenta",
        "descripcion": "No se pueden adjuntar documentos PDF de más de 5 MB.",
        "causa": "El límite de tamaño configurado era de 5 MB, menor al que requiere el proceso.",
        "solucion": "Se amplió el límite a 15 MB según lo acordado con el área usuaria.",
        "categoria": "Configuración_Adjuntos",
        "subcategoria": "Formatos y tamaño permitidos",
        "despliegue": False,
    },
    {
        "flujo": "Solicitud de crédito",
        "descripcion": "El sistema deja avanzar sin adjuntar el certificado laboral, que es obligatorio.",
        "causa": "La regla de obligatoriedad del adjunto no se migró al nuevo formulario.",
        "solucion": "Se configuró el adjunto como obligatorio en el formulario.",
        "categoria": "Configuración_Adjuntos",
        "subcategoria": "Obligatoriedad del adjunto",
        "despliegue": False,
    },
    {
        "flujo": "Reclamos",
        "descripcion": "Los adjuntos cargados no se pueden visualizar desde la bandeja del analista.",
        "causa": "El visor de documentos no tenía permiso de lectura sobre la carpeta de adjuntos del flujo.",
        "solucion": "Se otorgó permiso de lectura al visor sobre la carpeta de adjuntos.",
        "categoria": "Configuración_Adjuntos",
        "subcategoria": "Carga y visualización de adjuntos",
        "despliegue": False,
    },
    {
        "flujo": "Reclamos",
        "descripcion": "La plantilla descargable de la carta de respuesta tiene el logo anterior.",
        "causa": "La plantilla no se actualizó en la última renovación de imagen.",
        "solucion": "Se reemplazó la plantilla por la versión vigente.",
        "categoria": "Configuración_Adjuntos",
        "subcategoria": "Plantillas descargables",
        "despliegue": False,
    },
    {
        "flujo": "Gestión de garantías",
        "descripcion": "La carga masiva de garantías falla con archivos de más de 500 filas.",
        "causa": "El proceso de carga tenía un tiempo de espera insuficiente para archivos grandes.",
        "solucion": "Se cambió la carga para procesar el archivo por lotes de 200 filas.",
        "categoria": "Configuración_Adjuntos",
        "subcategoria": "Carga masiva",
        "despliegue": True,
    },
]

# Casos registrados sin detalle: con esto el modelo no tiene de dónde sacar la causa y debe
# marcar [SIN SUSTENTO] en lugar de inventarla.
DESCRIPCIONES_POBRES = [
    "No funciona.",
    "Error en el flujo, revisar.",
    "El usuario reporta falla.",
    "Ver adjunto.",
]

DESCRIPCIONES_EVOLUTIVO = [
    "Agregar el campo 'Correo alternativo' al formulario de radicación.",
    "Incluir un filtro por fecha de radicación en la bandeja de analistas.",
    "Crear un informe mensual de reclamos por canal.",
    "Permitir reasignar casos en bloque desde la bandeja del supervisor.",
    "Agregar la firma digital del analista a la carta de respuesta.",
]

# Las formas en que el equipo escribe a mano la columna Causa/Solucion. Varían a propósito:

FORMATOS_CAUSA_SOLUCION = [
    "Causa: {causa}\nSolucion: {solucion}",
    "Causa: {causa} Solución: {solucion}",
    "CAUSA - {causa}\nSOLUCION - {solucion}",
    "causa : {causa}",
    "Se revisó el caso. Causa: {causa} Solución: {solucion}",
    "Causa: {causa}\n\nSolucion: {solucion}",
    "Causa:{causa}\nSolución:{solucion}",
    "Solucion: {solucion}",
]

OBSERVACIONES_ADMINISTRATIVAS = [
    "Se asigna a {dev}.",
    "Se solicita información adicional al usuario.",
    "Desplegado en QA el {fecha}.",
    "Pendiente VoBo del usuario.",
    "Usuario confirma funcionamiento.",
    "Se escala a fábrica externa.",
    "Se reprograma por prioridad del sprint.",
]

DEVS = [f"Dev {n:02d}" for n in range(1, 6)]
QAS = [f"QA {n:02d}" for n in range(1, 4)]
USUARIOS = [f"Usuario {n:02d}" for n in range(1, 13)]


def fecha_texto(d):
    return d.strftime("%d/%m/%Y")


def numero_ha(rng, usados):
    while True:
        n = rng.randint(7_700_000, 7_799_999)
        if n not in usados:
            usados.add(n)
            return n


def observaciones(rng, dev, recepcion, pista=None):
    lineas = rng.sample(OBSERVACIONES_ADMINISTRATIVAS, k=rng.randint(1, 3))
    texto = [l.format(dev=dev, fecha=fecha_texto(recepcion + timedelta(days=rng.randint(3, 15)))) for l in lineas]
    if pista:
        texto.append(pista)
    return "\n".join(texto)


def construir_casos(rng):
    #Devuelve la lista de filas (una por caso) en el orden de ENCABEZADOS.
    usados = set()
    casos = []

    # Plan de los 60 casos: (tipo, grupo). Los cerrados son los que la plataforma documenta.
    plan = (
        [("Incidente", "cerrado")] * 20
        + [("Evento", "cerrado")] * 8
        + [("Control de cambio", "cerrado")] * 7
        + [(rng.choice(["Incidente", "Evento", "Control de cambio"]), "abierto") for _ in range(15)]
        + [("Evolutivo", "evolutivo")] * 10
    )

    # Entre los 35 cerrados: cuáles traen Causa/Solucion, cuáles tienen descripción pobre y cuáles traen la categoría diligenciada. Se eligen por posición para que no se crucen.
    indices_cerrados = [i for i, (_, g) in enumerate(plan) if g == "cerrado"]
    rng.shuffle(indices_cerrados)
    con_causa_solucion = set(indices_cerrados[:8])
    con_descripcion_pobre = set(indices_cerrados[8:12])
    con_categoria = set(indices_cerrados[12:15])

    for i, (tipo, grupo) in enumerate(plan):
        escenario = rng.choice(ESCENARIOS)
        dev = rng.choice(DEVS)
        recepcion = date(2026, 6, 1) + timedelta(days=rng.randint(0, 110))
        asignado = recepcion + timedelta(days=rng.randint(0, 2))
        compromiso_dev = asignado + timedelta(days=rng.randint(3, 8))
        fin_dev = compromiso_dev + timedelta(days=rng.randint(-2, 3))
        compromiso_qa = fin_dev + timedelta(days=rng.randint(1, 3))
        fin_qa = compromiso_qa + timedelta(days=rng.randint(0, 2))

        flujo = escenario["flujo"]
        descripcion = escenario["descripcion"]
        etapa, estado = rng.choice(ESTADOS_ABIERTOS)
        categoria = subcategoria = causa_solucion = None
        pista = None
        cierre = salida_prd = vobo_solicitud = vobo = None
        tiene_qa = True

        if grupo == "evolutivo":
            descripcion = rng.choice(DESCRIPCIONES_EVOLUTIVO)
            etapa, estado = None, "Evolutivo"
            tiene_qa = rng.random() < 0.5
        elif grupo == "cerrado":
         
            if not escenario["despliegue"] and rng.random() < 0.5:
                etapa, estado = CERRADO_SIN_PRD
            else:
                etapa, estado = CERRADO
            cierre = fin_qa + timedelta(days=rng.randint(1, 6))
            if estado == CERRADO[1]:
                salida_prd = cierre - timedelta(days=rng.randint(0, 2))

            if i in con_descripcion_pobre:
                descripcion = rng.choice(DESCRIPCIONES_POBRES)
            if i in con_causa_solucion:
                formato = FORMATOS_CAUSA_SOLUCION[len(casos) % len(FORMATOS_CAUSA_SOLUCION)]
                causa_solucion = formato.format(causa=escenario["causa"], solucion=escenario["solucion"])
            if i in con_categoria:
                categoria, subcategoria = escenario["categoria"], escenario["subcategoria"]
           
            if causa_solucion is None and i not in con_descripcion_pobre and rng.random() < 0.3:
                pista = "Se aplica ajuste: " + escenario["solucion"].split(" y ")[0].rstrip(".") + "."

        if tipo == "Control de cambio" and grupo != "evolutivo":
            vobo_solicitud = fin_qa + timedelta(days=1)
            vobo = vobo_solicitud + timedelta(days=rng.randint(0, 3)) if cierre else None

        ha = numero_ha(rng, usados)
        ha_dev = ha
        if rng.random() < 0.07:
            
            ha_dev = f"{ha}\nDUPLICADO CON HA-{numero_ha(rng, usados)}"

        dias_cierre = (cierre - recepcion).days if cierre else (date(2026, 10, 1) - recepcion).days

        casos.append([
            tipo,
            rng.choice(USUARIOS),
            recepcion,
            f"Sprint {14 + (recepcion - date(2026, 6, 1)).days // 14}",
            descripcion,
            FLUJOS[flujo],
            flujo,
            rng.choice(["Interna", "Externa"]),
            rng.choice(["Sí", "No"]),
            ha_dev,
            dev,
            asignado,
            compromiso_dev,
            fin_dev if grupo != "abierto" or rng.random() < 0.5 else None,
            rng.choice(["Sí", "No", None]),
            categoria,
            subcategoria,
            rng.choice(["Sí", "No", None]),
            etapa,
            estado,
            numero_ha(rng, usados) if tiene_qa else None,
            rng.choice(QAS) if tiene_qa else None,
            compromiso_qa if tiene_qa else None,
            fin_qa if tiene_qa and grupo == "cerrado" else None,
            vobo_solicitud,
            vobo,
            cierre,
            salida_prd,
            dias_cierre,
            observaciones(rng, dev, recepcion, pista),
            causa_solucion,
        ])

    rng.shuffle(casos)
    return casos


def escribir_excel(casos, salida):
    libro = Workbook()
    # Como en el archivo real, el seguimiento está en la hoja "Flujos" y no es la primera: así
    # se prueba que carga.py la busca por nombre en vez de leer la primera hoja.
    portada = libro.active
    portada.title = "Resumen"
    portada["A1"] = "Resumen del seguimiento — DATOS FICTICIOS"
    portada["A1"].font = Font(bold=True, size=14)
    hoja = libro.create_sheet(config.HOJA_SEGUIMIENTO)

    
    hoja["A1"] = "Seguimiento de casos — DATOS FICTICIOS"
    hoja["A1"].font = Font(bold=True, size=14)
    hoja["A2"] = "Generado por scripts/generar_ejemplo.py. Ningún dato proviene del archivo real."

    for columna, nombre in enumerate(ENCABEZADOS, start=1):
        celda = hoja.cell(row=config.FILA_ENCABEZADO, column=columna, value=nombre)
        celda.font = Font(bold=True)

    for fila, caso in enumerate(casos, start=config.FILA_ENCABEZADO + 1):
        for columna, valor in enumerate(caso, start=1):
            celda = hoja.cell(row=fila, column=columna, value=valor)
            if isinstance(valor, date):
                celda.number_format = "dd/mm/yyyy"
            if isinstance(valor, str) and "\n" in valor:
                celda.alignment = Alignment(wrap_text=True, vertical="top")

  
    libro.properties.created = libro.properties.modified = datetime(2026, 10, 20)

    salida.parent.mkdir(parents=True, exist_ok=True)
    libro.save(salida)


def main():
    rng = random.Random(SEMILLA)
    casos = construir_casos(rng)
    escribir_excel(casos, SALIDA)

    tipo, estado, causa_sol = 0, 19, 30
    cerrados = [c for c in casos if c[estado].split(" ")[0].rstrip(".") in config.ESTADOS_CERRADOS]
    documentables = [c for c in cerrados if c[tipo].lower() in config.TIPOS_DOCUMENTABLES]
    print(f"Archivo: {SALIDA}")
    print(f"Casos: {len(casos)} | cerrados documentables: {len(documentables)} | "
          f"con Causa/Solucion: {sum(1 for c in casos if c[causa_sol])}")


if __name__ == "__main__":
    main()
