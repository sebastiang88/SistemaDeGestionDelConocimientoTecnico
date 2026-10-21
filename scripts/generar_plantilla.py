#Genera plantillas/ficha.docx: la plantilla de Word con la que se arma cada ficha.
#Las marcas {{ ... }} y {%p ... %} son las que docxtpl reemplaza por los datos de la ficha

import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import config  

SALIDA = config.BASE_DIR / "plantillas" / "ficha.docx"
AZUL = RGBColor(0x1F, 0x4E, 0x79) 

def main():
    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(11)

    titulo = doc.add_heading("Ficha técnica", level=0)
    titulo.runs[0].font.color.rgb = AZUL
    doc.add_paragraph("HA {{ ha }} · {{ estado }}").runs[0].bold = True

    # Datos del caso en una tabla de dos columnas.
    datos = [
        ("Tipo", "{{ tipo }}"),
        ("Flujo", "{{ flujo }}"),
        ("Sprint", "{{ sprint }}"),
        ("Reportado por", "{{ usuario_reporta }}"),
        ("Responsable", "{{ responsable }}"),
        ("Fecha de recepción", "{{ fecha_recepcion }}"),
        ("Fecha de cierre", "{{ fecha_cierre }}"),
        ("Generada", "{{ generada }}"),
        ("Documentada en", "{{ documentada_en }}"),
        ("Motor", "{{ motor }}"),
    ]
    tabla = doc.add_table(rows=len(datos), cols=2)
    tabla.style = "Light List Accent 1"
    for fila, (etiqueta, marca) in zip(tabla.rows, datos):
        fila.cells[0].text = etiqueta
        fila.cells[0].paragraphs[0].runs[0].bold = True
        fila.cells[1].text = marca

    # Contenido de la ficha. {{r ... }} es texto enriquecido: respeta los saltos de línea.
    for titulo_seccion, marca in (
        ("Síntoma", "{{r sintoma }}"),
        ("Causa", "{{r causa }}"),
        ("Solución", "{{r solucion }}"),
        ("Categoría de la causa", "{{ categoria }}"),
    ):
        encabezado = doc.add_heading(titulo_seccion, level=1)
        encabezado.runs[0].font.color.rgb = AZUL
        doc.add_paragraph(marca)

    # Secciones opcionales: sólo aparecen si el desarrollador las llenó.
    doc.add_paragraph("{%p if lecciones %}")
    encabezado = doc.add_heading("Lección aprendida / observaciones", level=1)
    encabezado.runs[0].font.color.rgb = AZUL
    doc.add_paragraph("{{r lecciones }}")
    doc.add_paragraph("{%p endif %}")

    doc.add_paragraph("{%p if evidencias %}")
    encabezado = doc.add_heading("Evidencias", level=1)
    encabezado.runs[0].font.color.rgb = AZUL
    doc.add_paragraph("{%p for e in evidencias %}")
    doc.add_paragraph("{{ e.imagen }}")
    leyenda = doc.add_paragraph("{{ e.nombre }}")
    leyenda.runs[0].font.size = Pt(9)
    leyenda.runs[0].italic = True
    doc.add_paragraph("{%p endfor %}")
    doc.add_paragraph("{%p endif %}")

    # Preguntas y respuestas: sólo si la ficha tuvo preguntas. Las marcas {%p %} van solas en su párrafo; docxtpl repite o quita esos párrafos completos.
    doc.add_paragraph("{%p if preguntas %}")
    encabezado = doc.add_heading("Preguntas al responsable", level=1)
    encabezado.runs[0].font.color.rgb = AZUL
    doc.add_paragraph("{%p for p in preguntas %}")
    doc.add_paragraph("{{ p.orden }}. {{ p.texto }}").runs[0].bold = True
    doc.add_paragraph("{{r p.respuesta }}")
    doc.add_paragraph("{%p endfor %}")
    doc.add_paragraph("{%p endif %}")

    pie = doc.add_paragraph("Generada con SGCT — Sistema de Gestión del Conocimiento Técnico.")
    pie.runs[0].font.size = Pt(9)
    pie.runs[0].italic = True

    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    doc.save(SALIDA)
    print(f"Plantilla: {SALIDA}")

if __name__ == "__main__":
    main()